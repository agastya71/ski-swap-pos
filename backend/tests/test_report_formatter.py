from datetime import date, datetime, timezone

import pytest
from fastapi import HTTPException
from fastapi.responses import JSONResponse

from app.schemas.reports import (
    EndOfDayReport,
    SellerPayoutReport,
    SellerPayoutSaleLine,
    SellerPayoutSellerInfo,
    SellerPayoutUnsoldLine,
)


def _eod():
    return EndOfDayReport(
        event_id=1, event_name="Test Event",
        date_generated=date.today(),
        sales_count=5, voided_count=1,
        gross_revenue=100.00, mysl_total=30.00, seller_total=70.00,
        cash_total=50.00, check_total=30.00, cc_total=20.00,
        generated_at=datetime.now(timezone.utc),
    )


def _payout():
    return SellerPayoutReport(
        event_id=1, event_name="Test Event",
        seller_id=1, seller_code="ABC", seller_name="Jane Smith",
        seller_email=None,
        seller_info=SellerPayoutSellerInfo(
            seller_code="ABC", seller_name="Jane Smith", company=None, is_vendor=False,
            email=None, phone="612-555-0100", address=None, city=None, state=None, zip=None,
            commission_rate=0.30,
        ),
        items_consigned=2, items_sold=1, items_unsold=1, items_donated=0,
        gross_sales=20.00, mysl_total=6.00, seller_total=14.00,
        sales=[
            SellerPayoutSaleLine(item_code="ABC-001", description="Skis",
                                 date_of_sale=datetime.now(timezone.utc),
                                 quantity_sold=1.0, sell_price=20.00, extended_price=20.00,
                                 price_adjustment_reason="Price match",
                                 mysl_share=6.0, seller_share=14.0, commission_rate=0.30),
        ],
        unsold_items=[
            SellerPayoutUnsoldLine(item_code="ABC-002", description="Boots",
                                   quantity=1.0, remaining=1.0, price=10.00,
                                   status="available", donate_unsold=False,
                                   mysl_share=0.0, seller_share=0.0, commission_rate=0.30),
        ],
        generated_at=datetime.now(timezone.utc),
    )


def test_format_json_returns_json_response():
    from app.services.report_formatter import format_report
    resp = format_report(_eod(), "json", "eod_test")
    assert isinstance(resp, JSONResponse)


def test_format_csv_content_type():
    from app.services.report_formatter import format_report
    resp = format_report(_eod(), "csv", "eod_test")
    assert resp.media_type == "text/csv"
    assert b"gross_revenue" in resp.body


def test_format_csv_payout_has_sales_and_unsold_sections():
    from app.services.report_formatter import format_report
    resp = format_report(_payout(), "csv", "payout_test")
    assert b"ABC-001" in resp.body
    assert b"SALES" in resp.body
    assert b"UNSOLD" in resp.body
    assert b"612-555-0100" in resp.body  # seller contact info


def test_format_md_content_type():
    from app.services.report_formatter import format_report
    resp = format_report(_eod(), "md", "eod_test")
    assert resp.media_type == "text/markdown"
    assert b"End of Day" in resp.body


def test_format_pdf_content_type():
    from app.services.report_formatter import format_report
    resp = format_report(_eod(), "pdf", "eod_test")
    assert resp.media_type == "application/pdf"
    assert resp.body[:4] == b"%PDF"


def test_format_invalid_raises_422():
    from app.services.report_formatter import format_report
    with pytest.raises(HTTPException) as exc:
        format_report(_eod(), "xml", "eod_test")
    assert exc.value.status_code == 422


def test_payout_csv_includes_adjustment_reason_column():
    """The seller payout CSV SALES section carries the price-adjustment reason."""
    from app.services.report_formatter import format_report
    resp = format_report(_payout(), "csv", "payout")
    assert b"price_adjustment_reason" in resp.body
    assert b"Price match" in resp.body


def test_payout_md_includes_adjustment_reason_column():
    """The seller payout Markdown SALES table carries the reason."""
    from app.services.report_formatter import format_report
    resp = format_report(_payout(), "md", "payout")
    body = bytes(resp.body).decode()
    assert "Adj. Reason" in body
    assert "Price match" in body


def test_payout_xlsx_includes_adjustment_reason_column():
    """The per-seller XLSX workbook's Sales sheet carries the reason and keeps
    the totals row aligned under Extended / Due Seller."""
    import io

    import openpyxl

    from app.services.report_formatter import format_report
    payout = _payout()
    # openpyxl rejects tz-aware datetimes; live rows are stored naive.
    payout = payout.model_copy(update={
        "sales": [
            s.model_copy(update={
                "date_of_sale": s.date_of_sale.replace(tzinfo=None) if s.date_of_sale else None
            })
            for s in payout.sales
        ],
    })
    resp = format_report(payout, "xlsx", "payout")
    wb = openpyxl.load_workbook(io.BytesIO(resp.body))
    ws = wb["Sales"]
    headers = [c.value for c in ws[1]]
    assert "Adj. Reason" in headers
    reason_col = headers.index("Adj. Reason")
    assert ws.cell(row=2, column=reason_col + 1).value == "Price match"
    totals = [c.value for c in ws[ws.max_row]]
    assert totals[headers.index("Extended")] == 20.00
    assert totals[headers.index("Due Seller")] == 14.00


def test_format_unsold_report_donate_column():
    """The dedicated unsold report carries the per-item donate election."""
    import io

    from pypdf import PdfReader

    from app.schemas.reports import UnsoldItem, UnsoldItemsReport
    from app.services.report_formatter import format_report

    report = UnsoldItemsReport(
        event_id=1, event_name="Swap 2026",
        items=[
            UnsoldItem(seller_code="DU", seller_name="Donate Tester",
                       item_code="DU-001", description="Skis", category="Skis",
                       quantity=1.0, remaining=1.0, price=10.0, donate_unsold=True),
        ],
        total_items=1, total_value=10.0,
        generated_at=datetime.now(timezone.utc),
    )
    csv = format_report(report, "csv", "unsold_donate")
    assert b"donate" in csv.body
    assert b"Yes" in csv.body
    md = format_report(report, "md", "unsold_donate")
    assert b"| Donate |" in md.body
    pdf = format_report(report, "pdf", "unsold_donate")
    assert pdf.media_type == "application/pdf"
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.body)).pages)
    assert "Donate" in text and "Yes" in text


# ── Vendor equipment summary ─────────────────────────────────────────────────

def _vendor_summary():
    from app.schemas.reports import VendorCategoryLine, VendorEquipmentSummaryReport

    line = VendorCategoryLine(
        category="Skis", items_consigned=2, units_sold=2.0, gross_sales=150.0,
        mysl_share=45.0, seller_share=105.0, units_unsold=1.0, unsold_value=100.0,
    )
    total = VendorCategoryLine(
        category="TOTAL", items_consigned=2, units_sold=2.0, gross_sales=150.0,
        mysl_share=45.0, seller_share=105.0, units_unsold=1.0, unsold_value=100.0,
    )
    return VendorEquipmentSummaryReport(
        event_id=1, event_name="Test Event", seller_code="VEND1",
        seller_name="Summary Vendor", company="Summary Vendor",
        vendor_commission_rate=0.30, categories=[line], total=total,
        generated_at=datetime.now(timezone.utc),
    )


def test_vendor_summary_csv():
    from app.services.report_formatter import format_report
    resp = format_report(_vendor_summary(), "csv", "vendor_summary")
    assert b"category,items_consigned,units_sold" in resp.body
    assert b"Skis" in resp.body
    assert b"TOTAL" in resp.body


def test_vendor_summary_md():
    from app.services.report_formatter import format_report
    resp = format_report(_vendor_summary(), "md", "vendor_summary")
    body = bytes(resp.body).decode()
    assert "Vendor Equipment Summary" in body
    assert "Skis" in body
    assert "**TOTAL**" in body


def test_vendor_summary_xlsx():
    import io

    import openpyxl

    from app.services.report_formatter import format_report
    resp = format_report(_vendor_summary(), "xlsx", "vendor_summary")
    wb = openpyxl.load_workbook(io.BytesIO(resp.body))
    ws = wb.active
    assert ws is not None, "workbook has no active sheet"
    assert ws.cell(row=1, column=1).value == "Vendor Equipment Summary"
    headers = [c.value for c in ws[6]]  # row 5 is the blank separator
    assert headers[0] == "Category"
    assert ws.cell(row=7, column=1).value == "Skis"
    totals_row = [c.value for c in ws[8]]
    assert totals_row[0] == "TOTAL"
    assert totals_row[3] == 150.0  # gross


def test_vendor_summary_pdf():
    """PDF content is compressed — verify via pypdf text extraction."""
    import io

    from pypdf import PdfReader

    from app.services.report_formatter import format_report
    resp = format_report(_vendor_summary(), "pdf", "vendor_summary")
    assert resp.media_type == "application/pdf"
    reader = PdfReader(io.BytesIO(resp.body))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    assert "Skis" in text
    assert "TOTAL" in text
