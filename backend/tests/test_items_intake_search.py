"""Tests for the intake item search endpoints: GET /items/intake-search
(full-field search across all intake items for the active event) and
GET /items/intake-search/export (Excel export of the same results)."""

import io

import openpyxl
import pytest

from app.models.event import Event
from app.models.intake import Intake
from app.models.item import Item
from app.models.seller import Seller


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def search_data(db, active_event):
    """Two sellers with one intake each and two live items + one soft-deleted item."""
    s1 = Seller(event_id=active_event.id, code="A001", first_name="Jane", last_name="Smith")
    s2 = Seller(event_id=active_event.id, code="A002", company="Pioneer Sports", is_vendor=True)
    db.add_all([s1, s2])
    db.flush()
    i1 = Intake(seller_id=s1.id)
    i2 = Intake(seller_id=s2.id)
    db.add_all([i1, i2])
    db.flush()
    db.add_all([
        Item(intake_id=i1.id, seller_id=s1.id, code="A001-01", description="Red alpine skis",
             category="Skis", brand="Rossignol", type="Alpine Ski", color="Red", size="140cm",
             gender_age="Kids'", year=2018, price=75.0, quantity=2, remaining=1, status="sold"),
        Item(intake_id=i2.id, seller_id=s2.id, code="A002-01", description="Blue snowboard boots",
             category="Boots", brand="Atomic", type="Snowboard Boot", color="Blue",
             size="8", gender_age="Men's", year=2020, price=45.0, quantity=3, remaining=3, status="available"),
        Item(intake_id=i1.id, code="A001-02", description="Deleted green helmet", seller_id=s1.id,
             brand="Giro", price=25.0, is_deleted=True),
    ])
    db.commit()


def test_search_by_description(client, admin_token, search_data):
    resp = client.get("/items/intake-search", params={"q": "alpine"}, headers=_auth(admin_token))
    assert resp.status_code == 200
    data = resp.json()
    assert len(data) == 1
    assert data[0]["code"] == "A001-01"
    assert data[0]["seller_code"] == "A001"
    assert data[0]["seller_name"] == "Jane Smith"


def test_search_matches_brand_seller_code_and_vendor_company(client, admin_token, search_data):
    assert [r["code"] for r in client.get("/items/intake-search", params={"q": "rossignol"}, headers=_auth(admin_token)).json()] == ["A001-01"]
    assert [r["code"] for r in client.get("/items/intake-search", params={"q": "A002"}, headers=_auth(admin_token)).json()] == ["A002-01"]
    vendor = client.get("/items/intake-search", params={"q": "pioneer"}, headers=_auth(admin_token)).json()
    assert len(vendor) == 1 and vendor[0]["seller_name"] == "Pioneer Sports"


def test_search_matches_numeric_fields_as_text(client, admin_token, search_data):
    # Year and price are searched as text casts: "2018" hits the year column,
    # "45" hits price 45.0.
    assert [r["code"] for r in client.get("/items/intake-search", params={"q": "2018"}, headers=_auth(admin_token)).json()] == ["A001-01"]
    assert [r["code"] for r in client.get("/items/intake-search", params={"q": "45"}, headers=_auth(admin_token)).json()] == ["A002-01"]


def test_search_empty_q_returns_all_active_event_items(client, admin_token, search_data, active_event):
    resp = client.get("/items/intake-search", headers=_auth(admin_token))
    assert resp.status_code == 200
    codes = [r["code"] for r in resp.json()]
    assert "A001-01" in codes and "A002-01" in codes
    assert "A001-02" not in codes  # soft-deleted items are never returned


def test_search_status_filter(client, admin_token, search_data):
    resp = client.get("/items/intake-search", params={"q": "", "status": "sold"}, headers=_auth(admin_token))
    assert [r["code"] for r in resp.json()] == ["A001-01"]
    resp = client.get("/items/intake-search", params={"status": "available"}, headers=_auth(admin_token))
    assert [r["code"] for r in resp.json()] == ["A002-01"]


def test_search_excludes_other_event_items(client, db, admin_token, active_event, search_data):
    stale = Event(name="Old Swap 2024", year=2024, commission_rate=0.30, is_active=False)
    db.add(stale)
    db.flush()
    seller = Seller(event_id=stale.id, code="A001")
    db.add(seller)
    db.flush()
    intake = Intake(seller_id=seller.id)
    db.add(intake)
    db.flush()
    db.add(Item(intake_id=intake.id, seller_id=seller.id, code="OLD-01",
                description="Red alpine skis", brand="Rossignol", price=99.0))
    db.commit()
    resp = client.get("/items/intake-search", params={"q": "alpine"}, headers=_auth(admin_token))
    # Only the ACTIVE event's item matches (item codes are globally unique, so
    # the stale-event row uses its own code "OLD-01" with identical text fields).
    assert [r["code"] for r in resp.json()] == ["A001-01"]


def test_search_limit(client, admin_token, search_data):
    resp = client.get("/items/intake-search", params={"limit": 1}, headers=_auth(admin_token))
    assert len(resp.json()) == 1


def test_search_intake_role_allowed_cashier_forbidden(client, intake_token, cashier_token, search_data):
    assert client.get("/items/intake-search", headers=_auth(intake_token)).status_code == 200
    assert client.get("/items/intake-search", headers=_auth(cashier_token)).status_code == 403


def test_search_requires_auth(client, search_data):
    assert client.get("/items/intake-search").status_code == 403


def test_export_returns_xlsx_with_results(client, admin_token, search_data):
    resp = client.get("/items/intake-search/export", params={"q": "alpine"}, headers=_auth(admin_token))
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert resp.headers["content-disposition"].startswith("attachment; filename=intake-items-")
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None, "export workbook has no active sheet"
    assert ws.cell(row=1, column=1).value == "Code"
    assert ws.cell(row=1, column=17).value == "Seller Name"
    assert ws.cell(row=2, column=1).value == "A001-01"
    assert ws.cell(row=2, column=4).value == "Rossignol"
    assert ws.cell(row=2, column=16).value == "A001"
    assert ws.cell(row=2, column=17).value == "Jane Smith"


def test_export_requires_admin_or_intake(client, cashier_token, search_data):
    assert client.get("/items/intake-search/export", headers=_auth(cashier_token)).status_code == 403